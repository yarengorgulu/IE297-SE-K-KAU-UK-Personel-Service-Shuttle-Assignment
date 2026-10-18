"""Excel çalışma kitaplarını yalnızca okuma kipinde açan katman."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from excel_import.models import ExcelRow, SheetData, SheetInfo, WorkbookInfo
from normalization.text import is_blank, normalize_whitespace


SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}


class ExcelReadError(RuntimeError):
    """Kullanıcıya gösterilebilecek Excel okuma hatası."""


def _validate_path(path: str | Path) -> Path:
    resolved = Path(path).expanduser()
    if not resolved.is_file():
        raise ExcelReadError("Seçilen Excel dosyası bulunamadı.")
    if resolved.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ExcelReadError("Yalnızca .xlsx ve .xlsm Excel dosyaları desteklenir.")
    return resolved


def _has_value(values: Iterable[Any]) -> bool:
    return any(not is_blank(value) for value in values)


def _header_names(values: tuple[Any, ...]) -> tuple[str, ...]:
    headers: list[str] = []
    for index, value in enumerate(values, start=1):
        if is_blank(value):
            headers.append(f"SÜTUN_{index}")
        else:
            headers.append(normalize_whitespace(value))
    return tuple(headers)


def _open_workbook(path: Path):
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except PermissionError as error:
        raise ExcelReadError(
            "Excel dosyasına erişilemiyor. Dosya başka bir programda kilitli olabilir."
        ) from error
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as error:
        raise ExcelReadError(
            "Excel dosyası okunamadı. Dosya bozuk veya desteklenmeyen bir biçimde olabilir."
        ) from error
    except Exception as error:
        raise ExcelReadError(
            "Excel dosyası okunurken beklenmeyen bir hata oluştu."
        ) from error


def inspect_workbook(path: str | Path) -> WorkbookInfo:
    resolved = _validate_path(path)
    workbook = _open_workbook(resolved)
    sheets: list[SheetInfo] = []
    try:
        for worksheet in workbook.worksheets:
            info = SheetInfo(name=worksheet.title, is_empty=True)
            for row_number, values in enumerate(
                worksheet.iter_rows(values_only=True), start=1
            ):
                if _has_value(values):
                    info = SheetInfo(
                        name=worksheet.title,
                        is_empty=False,
                        header_row_number=row_number,
                        headers=_header_names(tuple(values)),
                    )
                    break
            sheets.append(info)
    finally:
        workbook.close()
    return WorkbookInfo(path=resolved, sheets=tuple(sheets))


def read_sheet(path: str | Path, sheet_name: str) -> SheetData:
    resolved = _validate_path(path)
    workbook = _open_workbook(resolved)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ExcelReadError(f"'{sheet_name}' adlı çalışma sayfası bulunamadı.")
        worksheet = workbook[sheet_name]
        header_row_number: int | None = None
        headers: tuple[str, ...] = ()
        rows: list[ExcelRow] = []
        skipped_blank_rows = 0

        for row_number, values in enumerate(
            worksheet.iter_rows(values_only=True), start=1
        ):
            values_tuple = tuple(values)
            if header_row_number is None:
                if not _has_value(values_tuple):
                    continue
                header_row_number = row_number
                headers = _header_names(values_tuple)
                continue

            relevant_values = values_tuple[: len(headers)]
            if not _has_value(relevant_values):
                skipped_blank_rows += 1
                continue
            if len(relevant_values) < len(headers):
                relevant_values += (None,) * (len(headers) - len(relevant_values))
            rows.append(ExcelRow(row_number=row_number, values=relevant_values))

        if header_row_number is None:
            raise ExcelReadError("Seçilen çalışma sayfası tamamen boş.")
        return SheetData(
            name=sheet_name,
            header_row_number=header_row_number,
            headers=headers,
            rows=tuple(rows),
            skipped_blank_rows=skipped_blank_rows,
        )
    finally:
        workbook.close()
